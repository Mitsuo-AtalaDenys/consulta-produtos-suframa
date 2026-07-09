#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
comparar_bases.py
=================
Compara duas bases da SUFRAMA e gera um relatorio Excel com:
  - NOVOS      : produto/tipo que existia so na base nova
  - REMOVIDOS  : produto/tipo que existia so na base anterior
  - ALTERADOS  : mesmo produto/tipo, mas com mudanca em algum campo
                 (NCM, descricao NCM, descricao do tipo, base legal, unidade)

Uso:

  # 1) ANTES de rodar a nova coleta, marque a base atual como "anterior":
  python comparar_bases.py --marcar-anterior

  # 2) Rode a coleta normalmente (gera nova suframa_insumos.parquet):
  python coletar_suframa.py --reiniciar
  # e depois:
  copy dados\\suframa_insumos.parquet suframa_insumos.parquet

  # 3) Gere o relatorio de mudancas:
  python comparar_bases.py --gerar-relatorio

Saida:
  - suframa_anterior.parquet  (copia da base anterior)
  - relatorio_atualizacao_YYYY-MM-DD.xlsx

Requisitos:
  pip install pandas pyarrow openpyxl
"""

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ATUAL = Path("suframa_insumos.parquet")
ANTERIOR = Path("suframa_anterior.parquet")

# Colunas exibidas no relatorio (na ordem)
COLS_EXIBICAO = [
    ("codigo_produto", "Código"),
    ("produto_nome", "Produto"),
    ("codigo_tipo", "Tipo"),
    ("descricao_tipo", "Descrição do Tipo"),
    ("ncm", "NCM"),
    ("ncm_descricao", "Descrição da NCM"),
    ("base_legal", "Base Legal"),
    ("unidade", "Unidade"),
    ("data_atualizacao", "Atualização"),
]

# Colunas comparadas para detectar ALTERACAO
COLS_COMPARADAS = [
    "descricao_tipo",
    "ncm",
    "ncm_descricao",
    "base_legal",
    "unidade",
]


# --- estilo Atala & Denys -------------------------------------------------
BURGUNDY = "3D1E24"
SAND = "DFD4C4"
CREAM = "F3F0EA"
WHITE = "FFFFFF"

FONT_HEADER = Font(name="Arial", size=11, bold=True, color=WHITE)
FONT_BODY = Font(name="Arial", size=10, color="1F1F1F")
FILL_HEADER = PatternFill("solid", start_color=BURGUNDY, end_color=BURGUNDY)
FILL_ROW = PatternFill("solid", start_color=CREAM, end_color=CREAM)
BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)


# --------------------------------------------------------------------------
def marcar_anterior():
    if not ATUAL.exists():
        sys.exit(f"[erro] {ATUAL} nao encontrado nesta pasta.")
    shutil.copy2(ATUAL, ANTERIOR)
    tam = ANTERIOR.stat().st_size / 1024
    print(f">> Base atual copiada para {ANTERIOR.name} ({tam:.1f} KB)")
    print("   Agora rode a coleta e depois: python comparar_bases.py --gerar-relatorio")


def carregar(caminho: Path) -> pd.DataFrame:
    df = pd.read_parquet(caminho).fillna("")
    for c, _ in COLS_EXIBICAO:
        if c not in df.columns:
            df[c] = ""
    df["_chave"] = (
        df["codigo_produto"].astype(str) + "|" + df["codigo_tipo"].astype(str)
    )
    return df


def comparar(anterior: pd.DataFrame, nova: pd.DataFrame) -> pd.DataFrame:
    chaves_ant = set(anterior["_chave"])
    chaves_nov = set(nova["_chave"])

    novos = nova[nova["_chave"].isin(chaves_nov - chaves_ant)].copy()
    novos["status"] = "NOVO"
    novos["campos_alterados"] = ""

    removidos = anterior[anterior["_chave"].isin(chaves_ant - chaves_nov)].copy()
    removidos["status"] = "REMOVIDO"
    removidos["campos_alterados"] = ""

    # ALTERADOS: mesmo (produto, tipo) mas algum campo mudou
    comuns = chaves_ant & chaves_nov
    ant_i = anterior.set_index("_chave")
    nov_i = nova.set_index("_chave")

    linhas_alt = []
    for k in comuns:
        a = ant_i.loc[k]
        n = nov_i.loc[k]
        mudou = [c for c in COLS_COMPARADAS if str(a.get(c, "")) != str(n.get(c, ""))]
        if mudou:
            linha = n.to_dict()
            linha["_chave"] = k
            linha["status"] = "ALTERADO"
            linha["campos_alterados"] = ", ".join(mudou)
            linhas_alt.append(linha)
    alterados = pd.DataFrame(linhas_alt) if linhas_alt else pd.DataFrame(columns=nova.columns.tolist() + ["status", "campos_alterados"])

    total = pd.concat([novos, removidos, alterados], ignore_index=True)
    return total


def formatar_planilha(ws, colunas, largura_por_coluna):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, largura in enumerate(largura_por_coluna, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.row_dimensions[1].height = 26

    n_linhas = ws.max_row
    for row in ws.iter_rows(min_row=1, max_row=n_linhas, max_col=len(colunas)):
        for cell in row:
            cell.border = BORDER
            if cell.row == 1:
                cell.font = FONT_HEADER
                cell.fill = FILL_HEADER
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = FONT_BODY
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def escrever_aba(wb: Workbook, nome: str, df: pd.DataFrame, incluir_campos_alterados: bool):
    ws = wb.create_sheet(nome)
    cols = [("status", "Status")] + COLS_EXIBICAO
    if incluir_campos_alterados:
        cols.append(("campos_alterados", "Campos Alterados"))

    # Cabecalho
    ws.append([rot for _, rot in cols])

    if df.empty:
        ws.append(["(nenhum registro nesta categoria)"] + [""] * (len(cols) - 1))
    else:
        for _, r in df.iterrows():
            ws.append([str(r.get(c, "")) for c, _ in cols])

    larguras = {
        "status": 12, "codigo_produto": 10, "produto_nome": 45,
        "codigo_tipo": 8, "descricao_tipo": 35, "ncm": 12,
        "ncm_descricao": 40, "base_legal": 40, "unidade": 12,
        "data_atualizacao": 14, "campos_alterados": 30,
    }
    formatar_planilha(ws, cols, [larguras[c] for c, _ in cols])


def escrever_resumo(wb: Workbook, ant: pd.DataFrame, nov: pd.DataFrame,
                    n_novos: int, n_removidos: int, n_alterados: int):
    ws = wb.create_sheet("Resumo", 0)

    dados = [
        ("RELATÓRIO DE ATUALIZAÇÃO — BASE SUFRAMA", ""),
        ("", ""),
        ("Data do relatório", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        ("Base anterior — total de linhas (tipos)", len(ant)),
        ("Base anterior — produtos únicos", ant["codigo_produto"].nunique()),
        ("Base atual — total de linhas (tipos)", len(nov)),
        ("Base atual — produtos únicos", nov["codigo_produto"].nunique()),
        ("", ""),
        ("Registros NOVOS", n_novos),
        ("Registros REMOVIDOS", n_removidos),
        ("Registros ALTERADOS", n_alterados),
        ("Total de mudanças", n_novos + n_removidos + n_alterados),
    ]
    for row in dados:
        ws.append(list(row))

    # estilo do resumo
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=BURGUNDY)
    for r in range(3, len(dados) + 1):
        ws[f"A{r}"].font = Font(name="Arial", size=11, bold=True)
        ws[f"B{r}"].font = Font(name="Arial", size=11)
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 22


def gerar_relatorio():
    if not ANTERIOR.exists():
        sys.exit(f"[erro] {ANTERIOR.name} nao encontrado. "
                 "Rode antes: python comparar_bases.py --marcar-anterior")
    if not ATUAL.exists():
        sys.exit(f"[erro] {ATUAL.name} nao encontrado.")

    print(">> Carregando bases...")
    ant = carregar(ANTERIOR)
    nov = carregar(ATUAL)
    print(f"   Anterior : {len(ant):>5} linhas · {ant['codigo_produto'].nunique()} produtos")
    print(f"   Atual    : {len(nov):>5} linhas · {nov['codigo_produto'].nunique()} produtos")

    print(">> Comparando...")
    mudancas = comparar(ant, nov)
    n_novos = int((mudancas["status"] == "NOVO").sum())
    n_rem = int((mudancas["status"] == "REMOVIDO").sum())
    n_alt = int((mudancas["status"] == "ALTERADO").sum())
    print(f"   NOVOS     : {n_novos}")
    print(f"   REMOVIDOS : {n_rem}")
    print(f"   ALTERADOS : {n_alt}")

    print(">> Montando planilha...")
    wb = Workbook()
    wb.remove(wb.active)  # remove a aba default vazia

    escrever_resumo(wb, ant, nov, n_novos, n_rem, n_alt)
    escrever_aba(wb, "NOVOS", mudancas[mudancas["status"] == "NOVO"], False)
    escrever_aba(wb, "REMOVIDOS", mudancas[mudancas["status"] == "REMOVIDO"], False)
    escrever_aba(wb, "ALTERADOS", mudancas[mudancas["status"] == "ALTERADO"], True)

    data = datetime.now().strftime("%Y-%m-%d")
    saida = Path(f"relatorio_atualizacao_{data}.xlsx")
    wb.save(saida)
    print(f"\n>> Concluido! Relatorio salvo em: {saida.resolve()}")


def main():
    p = argparse.ArgumentParser(description="Compara duas coletas da SUFRAMA e gera relatorio Excel")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--marcar-anterior", action="store_true",
                   help="copia a base atual como referencia para comparar depois")
    g.add_argument("--gerar-relatorio", action="store_true",
                   help="compara a base anterior com a atual e gera o Excel")
    a = p.parse_args()

    if a.marcar_anterior:
        marcar_anterior()
    else:
        gerar_relatorio()


if __name__ == "__main__":
    main()
